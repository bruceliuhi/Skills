"""
Excel 公式重算与错误检查
通过 LibreOffice 的独立配置副本转换计算公式，保留原始文件。
"""

import argparse
import contextlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from soffice import get_soffice_env, run_soffice

try:
    from openpyxl import load_workbook
except ImportError:
    print(json.dumps({"error": "当前解释器没有 openpyxl，未进行重算。请先运行 runtime_check.py，使用已验证的解释器或原生表格工具。"}, ensure_ascii=False))
    sys.exit(2)

SOFFICE_MISSING = "当前环境没有可执行的 LibreOffice（soffice）；公式尚未重算。可改用当前会话已提供的原生表格计算工具，并重新读取实际输出文件验收。"
MAX_LOCATIONS = 100
EXTERNAL_REF_RE = re.compile(r"""(?<![\w"\[])'?\[\d+\][^!"\[\]]*'?!""")


def _protected_features(filename):
    with zipfile.ZipFile(filename) as archive:
        names = archive.namelist()
        unsafe_formula = re.compile(r"(?i)(?:^|[^A-Z0-9_])(?:_xlfn\.)?(?:WEBSERVICE|RTD|DDE|STOCKHISTORY|IMAGE|SQLREQUEST|EVALUATE)\s*\(")
        blocked_calls = False
        for member in names:
            if member == "xl/workbook.xml" or (member.startswith("xl/worksheets/") and member.endswith(".xml")):
                root = ET.fromstring(archive.read(member))
                for element in root.iter():
                    if element.tag.rsplit("}", 1)[-1] in ("f", "definedName") and unsafe_formula.search(element.text or ""):
                        blocked_calls = True
    prefixes = {
        "xl/externalLinks/": "外部工作簿链接",
        "xl/pivotTables/": "透视表",
        "xl/pivotCache/": "透视缓存",
        "xl/slicers/": "切片器",
        "xl/slicerCaches/": "切片器缓存",
        "xl/activeX/": "ActiveX 控件",
        "xl/ctrlProps/": "表单控件",
        "xl/embeddings/": "嵌入对象",
        "xl/model/": "数据模型",
    }
    features = [label for prefix, label in prefixes.items() if any(n.startswith(prefix) for n in names)]
    if blocked_calls:
        features.append("外部通信或程序调用公式（需单独确认）")
    if "xl/vbaProject.bin" in names:
        features.append("VBA 宏")
    if "xl/connections.xml" in names:
        features.append("外部数据连接")
    return features


def external_links_at_risk(filename):
    try:
        with zipfile.ZipFile(filename) as archive:
            names = archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return []
    if not any(n.startswith("xl/externalLinks/") for n in names):
        return []

    with contextlib.ExitStack() as stack:
        formulas = load_workbook(filename, data_only=False)
        stack.callback(formulas.close)
        values = load_workbook(filename, data_only=True)
        stack.callback(values.close)

        external_names = [
            name
            for name, dn in formulas.defined_names.items()
            if isinstance(getattr(dn, "value", None), str) and EXTERNAL_REF_RE.search(dn.value)
        ]
        name_re = (
            re.compile(r"\b(" + "|".join(re.escape(n) for n in external_names) + r")\b")
            if external_names
            else None
        )

        at_risk = []
        for sheet in formulas.sheetnames:
            ws = formulas[sheet]
            if not hasattr(ws, "iter_rows"):  
                continue
            cached = values[sheet]
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    reaches_out = EXTERNAL_REF_RE.search(v) or (name_re and name_re.search(v))
                    if reaches_out and cached[cell.coordinate].value is None:
                        at_risk.append(f"{sheet}!{cell.coordinate}")
        return at_risk


def _font_family_changes(original, converted):
    """Reject known conversion loss before replacing the output copy."""
    before = load_workbook(original, data_only=False)
    after = load_workbook(converted, data_only=False)
    changes = []
    try:
        if before.sheetnames != after.sheetnames:
            return [{"reason": "工作表名称或顺序改变"}]
        for name in before.sheetnames:
            for row in before[name].iter_rows():
                for cell in row:
                    if cell.value is None or cell.font.name is None:
                        continue
                    actual = after[name][cell.coordinate].font.name
                    if actual != cell.font.name:
                        changes.append({"sheet": name, "cell": cell.coordinate,
                                        "original_font": cell.font.name,
                                        "converted_font": actual})
        return changes
    finally:
        before.close()
        after.close()


def recalc(filename, timeout=30, force=False):
    """Internal operation on an already-created output copy, never the source file."""
    path = Path(filename)
    if not path.is_file():
        return {"error": f"无法读取工作簿：{filename}"}
    if path.suffix.lower() != ".xlsx":
        return {"error": "自动重算仅支持普通 .xlsx；宏、模板或旧格式需使用已验证的原生 Office 路径。"}
    if not isinstance(timeout, int) or isinstance(timeout, bool) or timeout < 1 or timeout > 600:
        return {"error": "重算超时秒数必须是 1–600 的整数。"}
    if not os.access(path, os.W_OK):
        return {"error": "目标副本不可写；未进行公式重算。"}
    try:
        features = _protected_features(path)
        if features:
            return {"error": "此文件含尚未验证保真的复杂对象，已停止自动转换；请使用原生 Office 对副本处理并逐项验证。", "protected_features": features}
        # Keep the source's external-link inspection, and never force a destructive save.
        at_risk = external_links_at_risk(path)
        if at_risk:
            return {"error": "外部链接的缓存缺失，已停止重算以保护原有链接。", "external_link_cells": at_risk[:MAX_LOCATIONS]}
        get_soffice_env()
    except Exception as error:
        return {"error": f"无法准备公式重算：{error}"}
    with tempfile.TemporaryDirectory(prefix="excel-recalc-") as temporary:
        return _recalc_with_profile(filename, str(path.resolve()), timeout, Path(temporary))


def _recalc_with_profile(filename, abs_path, timeout, profile_dir):
    # Normal LibreOffice CLI conversion performs calculation while saving to a distinct directory.
    # It avoids installation or invocation of a Basic macro, which timed out on the tested macOS build.
    input_dir = profile_dir / "input"
    output_dir = profile_dir / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    input_copy = input_dir / Path(filename).name
    shutil.copyfile(abs_path, input_copy)
    # Force a full calculation on open through workbook metadata; preserve all other ZIP members.
    with zipfile.ZipFile(input_copy) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    rewritten = []
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    for item, data in members:
        if item.filename == "xl/workbook.xml":
            # Modify only calcPr in the private conversion copy, without serializing unrelated XML.
            text = data.decode("utf-8")
            root_match = re.search(r"<((?:[A-Za-z_][\w.-]*:)?)workbook\b", text)
            if root_match is None:
                raise ValueError("工作簿主 XML 格式不可识别，未进行转换。")
            prefix = root_match.group(1)
            calc = f'<{prefix}calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
            if re.search(r"<(?:[A-Za-z_][\w.-]*:)?calcPr\b", text):
                text = re.sub(r"<(?:[A-Za-z_][\w.-]*:)?calcPr\b[^>]*(?:/>|>.*?</(?:[A-Za-z_][\w.-]*:)?calcPr>)", calc, text, count=1, flags=re.S)
            else:
                text = text.replace(f"</{prefix}workbook>", calc + f"</{prefix}workbook>")
            data = text.encode("utf-8")
        rewritten.append((item, data))
    with zipfile.ZipFile(input_copy, "w") as archive:
        for item, data in rewritten:
            archive.writestr(item, data)
    args = [
        "--headless", "--norestore",
        f"-env:UserInstallation={(profile_dir / 'profile').as_uri()}",
        "--convert-to", "xlsx:Calc MS Excel 2007 XML",
        "--outdir", str(output_dir), str(input_copy),
    ]
    try:
        completed = run_soffice(args, capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return {"error": f"LibreOffice 在 {timeout} 秒内未完成；公式尚未验收，目标副本未替换。"}
    except FileNotFoundError:
        return {"error": SOFFICE_MISSING}
    except Exception as error:
        return {"error": f"公式重算未完成：{error}"}
    converted = output_dir / input_copy.name
    if completed.returncode != 0 or not converted.is_file():
        return {"error": "LibreOffice 未生成可验收的工作簿，公式尚未确认重算。", "engine_exit_code": completed.returncode, "engine_message": (completed.stderr or completed.stdout).strip()}
    try:
        with zipfile.ZipFile(converted) as archive:
            if archive.testzip() is not None:
                return {"error": "重算输出压缩包校验失败；目标副本未替换。"}
        check = load_workbook(converted, data_only=False)
        check.close()
        font_changes = _font_family_changes(abs_path, converted)
        if font_changes:
            return {"error": "计算引擎改变了原有字体，已拒绝替换目标副本。请使用已验证可保存的原生表格路径，并重新核验公式和样式。",
                    "preservation_issue": "font_family_changed",
                    "font_changes": font_changes[:MAX_LOCATIONS],
                    "font_change_count": len(font_changes)}
        shutil.copyfile(converted, abs_path)
    except Exception as error:
        return {"error": f"重算输出不能正常重新读取：{error}"}

    try:
        wb = load_workbook(filename, data_only=True)

        excel_errors = [
            "#VALUE!",
            "#DIV/0!",
            "#REF!",
            "#NAME?",
            "#NULL!",
            "#NUM!",
            "#N/A",
            "#SPILL!",
            "#CALC!",
            "#GETTING_DATA",
        ]
        error_details = {err: [] for err in excel_errors}
        total_errors = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if not hasattr(ws, "iter_rows"):  
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "e" and isinstance(cell.value, str):
                        for err in excel_errors:
                            if err == cell.value:
                                location = f"{sheet_name}!{cell.coordinate}"
                                error_details[err].append(location)
                                total_errors += 1
                                break

        result = {
            "status": "success" if total_errors == 0 else "errors_found",
            "total_errors": total_errors,
            "error_summary": {},
        }

        for err_type, locations in error_details.items():
            if locations:
                entry = {"count": len(locations), "locations": locations[:MAX_LOCATIONS]}
                if len(locations) > MAX_LOCATIONS:
                    entry["locations_truncated"] = len(locations) - MAX_LOCATIONS
                result["error_summary"][err_type] = entry

        wb.close()

        wb_formulas = load_workbook(filename, data_only=False)
        formula_count = 0
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            if not hasattr(ws, "iter_rows"):  
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if (
                        cell.value
                        and isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    ):
                        formula_count += 1
        wb_formulas.close()

        result["total_formulas"] = formula_count
        result["calculation_method"] = "libreoffice_headless_copy_conversion"
        result["status_zh"] = "公式重算未发现错误值" if total_errors == 0 else "发现公式错误，需修正"

        return result

    except Exception as e:
        return {"error": f"重算后的工作簿校验失败：{e}"}


class ChineseParser(argparse.ArgumentParser):
    def error(self, message):
        for old, new in [("the following arguments are required:", "缺少必填参数："), ("unrecognized arguments:", "无法识别的参数："), ("argument --timeout: invalid int value:", "--timeout 必须为整数："), ("expected one argument", "需要一个参数值")]:
            message = message.replace(old, new)
        print("输入有误：" + message, file=sys.stderr)
        print("用法：python recalc.py 输入.xlsx --output 新成果.xlsx [--timeout 秒数]", file=sys.stderr)
        self.exit(2)


def main():
    parser = ChineseParser(description="在新副本中重算 Excel 公式；保留原始文件，不覆盖已有成果。", add_help=False)
    parser.add_argument("-h", "--help", action="help", help="显示帮助并退出")
    parser._positionals.title = "位置参数"
    parser._optionals.title = "选项"
    parser.add_argument("excel_file", help="用户指定的原始 .xlsx 路径")
    parser.add_argument("--output", required=True, help="尚不存在的新 .xlsx 输出路径")
    parser.add_argument("--timeout", type=int, default=30, help="重算超时秒数（1–600，默认 30）")
    args = parser.parse_args()
    source = Path(args.excel_file).expanduser().resolve()
    destination = Path(args.output).expanduser().resolve()
    if not source.is_file():
        parser.error("无法读取输入文件；请确认当前上传或明确指定的文件路径。")
    if source == destination:
        parser.error("输出必须是新副本，不能与原始文件相同。")
    if destination.exists():
        parser.error("输出文件已存在；请换一个新文件名，避免覆盖既有成果。")
    if source.suffix.lower() != ".xlsx" or destination.suffix.lower() != ".xlsx":
        parser.error("此安全重算入口仅支持 .xlsx 输入与输出。")
    if not 1 <= args.timeout <= 600:
        parser.error("--timeout 必须为 1–600 秒。")
    try:
        with source.open("rb") as src, destination.open("xb") as dst:
            shutil.copyfileobj(src, dst)
        result = recalc(str(destination), args.timeout)
    except (OSError, ValueError) as error:
        result = {"error": f"不能创建或处理目标副本：{error}"}
    result["source_file"] = str(source)
    result["output_file"] = str(destination) if destination.is_file() else None
    if "error" in result:
        result["output_state_zh"] = "未完成验收，不可作为已重算成果交付"
    print(json.dumps(result, ensure_ascii=False, indent=2))
    sys.exit(1 if "error" in result else (2 if result["status"] == "errors_found" else 0))


if __name__ == "__main__":
    main()
